"""Generate samples/equipment-rental-booking.docx (a DOCX sample for the upload pipeline).

Run: python scripts/build_sample_docx.py
"""

from __future__ import annotations

from pathlib import Path

import docx

OUTPUT = Path(__file__).resolve().parents[1] / "samples" / "equipment-rental-booking.docx"


def build() -> Path:
    document = docx.Document()
    document.add_heading("RentRight Equipment Booking - Requirements", level=0)
    document.add_paragraph(
        "Fictional sample document for the ai-requirements-agent demo. RentRight rents "
        "construction equipment (excavators, scissor lifts, generators) from regional depots."
    )

    document.add_heading("Functional requirements", level=1)
    for line in [
        "Customers shall be able to search equipment by category, depot and rental dates.",
        "The system shall allow customers to book equipment for a date range and pay a deposit.",
        "Each booking has a booking number, customer, equipment item, start date, end date, "
        "deposit amount and status (requested, confirmed, out, returned, cancelled).",
        "Each equipment item has a serial number, category, depot, daily rate and "
        "maintenance status.",
        "Depot staff must be able to confirm bookings and record equipment hand-over.",
        "Customers can cancel bookings up to 48 hours before the start date.",
        "When equipment is returned late, the system shall charge the daily rate for each "
        "extra day.",
    ]:
        document.add_paragraph(line, style="List Bullet")

    document.add_heading("Non-functional requirements", level=1)
    for line in [
        "Availability search shall return results within 1 second for 95% of requests.",
        "Payment card data must never be stored by the system; payments use a PCI-compliant "
        "payment provider.",
        "The booking site should meet WCAG 2.1 AA.",
    ]:
        document.add_paragraph(line, style="List Bullet")

    table = document.add_table(rows=3, cols=2)
    for row, (term, meaning) in enumerate(
        [
            ("Term", "Meaning"),
            ("Depot", "Physical branch holding equipment"),
            ("Hand-over", "Equipment leaving the depot with the customer"),
        ]
    ):
        table.cell(row, 0).text, table.cell(row, 1).text = term, meaning

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    document.save(OUTPUT)
    return OUTPUT


if __name__ == "__main__":
    print(f"wrote {build()}")
