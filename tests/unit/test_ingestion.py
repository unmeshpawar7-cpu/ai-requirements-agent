from __future__ import annotations

import io
from itertools import pairwise

import docx
import pytest

from requirements_agent.ingestion.chunker import Chunker, detect_heading
from requirements_agent.ingestion.extraction import (
    DocumentFormat,
    ExtractionError,
    UnsupportedFormatError,
    detect_format,
    extract_text,
)
from requirements_agent.ingestion.normalise import normalise_text
from requirements_agent.ingestion.tokens import estimate_tokens


def make_pdf(lines: list[str]) -> bytes:
    """Build a minimal, valid single-page PDF with Helvetica text (no external tools)."""
    content = (
        "BT /F1 12 Tf 72 720 Td 14 TL "
        + " ".join(f"({line.replace('(', '[').replace(')', ']')}) '" for line in lines)
        + " ET"
    )
    objects = [
        "<< /Type /Catalog /Pages 2 0 R >>",
        "<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        "<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents 4 0 R "
        "/Resources << /Font << /F1 5 0 R >> >> >>",
        f"<< /Length {len(content)} >>\nstream\n{content}\nendstream",
        "<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    out = io.BytesIO()
    out.write(b"%PDF-1.4\n")
    offsets = []
    for number, body in enumerate(objects, start=1):
        offsets.append(out.tell())
        out.write(f"{number} 0 obj\n{body}\nendobj\n".encode("latin-1"))
    xref = out.tell()
    out.write(f"xref\n0 {len(objects) + 1}\n0000000000 65535 f \n".encode())
    for offset in offsets:
        out.write(f"{offset:010d} 00000 n \n".encode())
    out.write(
        f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode()
    )
    return out.getvalue()


# --------------------------------------------------------------------------- extraction
@pytest.mark.parametrize(
    ("filename", "content_type", "expected"),
    [
        ("spec.md", None, DocumentFormat.MARKDOWN),
        ("SPEC.TXT", None, DocumentFormat.TEXT),
        ("a.pdf", "application/octet-stream", DocumentFormat.PDF),
        ("a.docx", None, DocumentFormat.DOCX),
        (None, "text/markdown; charset=utf-8", DocumentFormat.MARKDOWN),
        ("noext", "application/pdf", DocumentFormat.PDF),
    ],
)
def test_detect_format(filename, content_type, expected) -> None:
    assert detect_format(filename, content_type) is expected


def test_detect_format_rejects_unknown_types() -> None:
    with pytest.raises(UnsupportedFormatError):
        detect_format("payload.exe", "application/octet-stream")


def test_text_decoding_handles_bom_utf16_and_legacy_encodings() -> None:
    assert extract_text("\ufeffShall work".encode(), DocumentFormat.TEXT).text == "Shall work"
    assert (
        extract_text("R\u00e9sum\u00e9".encode("utf-16"), DocumentFormat.TEXT).text
        == "R\u00e9sum\u00e9"
    )
    assert extract_text("caf\u00e9".encode("cp1252"), DocumentFormat.TEXT).text == "caf\u00e9"


def test_text_extraction_rejects_binary_and_empty() -> None:
    with pytest.raises(ExtractionError):
        extract_text(b"\x00\x01binary", DocumentFormat.TEXT)
    with pytest.raises(ExtractionError):
        extract_text(b"", DocumentFormat.MARKDOWN)


def test_pdf_extraction() -> None:
    data = make_pdf(["The system shall export invoices.", "Approvers must sign in."])
    result = extract_text(data, DocumentFormat.PDF)
    assert "export invoices" in result.text
    assert "Approvers must sign in" in result.text
    assert result.page_count == 1


def test_pdf_extraction_errors() -> None:
    with pytest.raises(ExtractionError, match="does not look like a PDF"):
        extract_text(b"hello", DocumentFormat.PDF)
    with pytest.raises(ExtractionError):
        extract_text(b"%PDF-1.4 garbage", DocumentFormat.PDF)
    with pytest.raises(ExtractionError, match="no extractable text"):
        extract_text(make_pdf([]), DocumentFormat.PDF)


def test_docx_extraction_maps_headings_lists_and_tables() -> None:
    document = docx.Document()
    document.add_heading("Invoice Portal", level=0)
    document.add_heading("Approvals", level=1)
    document.add_paragraph("Approvers must approve invoices.")
    document.add_paragraph("Escalate after 5 days", style="List Bullet")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text, table.cell(0, 1).text = "Field", "Type"
    table.cell(1, 0).text, table.cell(1, 1).text = "amount", "decimal"
    buffer = io.BytesIO()
    document.save(buffer)

    text = extract_text(buffer.getvalue(), DocumentFormat.DOCX).text
    assert "# Invoice Portal" in text
    assert "## Approvals" in text
    assert "- Escalate after 5 days" in text
    assert "amount | decimal" in text


def test_docx_extraction_errors() -> None:
    with pytest.raises(ExtractionError):
        extract_text(b"not a zip", DocumentFormat.DOCX)
    with pytest.raises(ExtractionError):
        extract_text(b"PK\x03\x04broken", DocumentFormat.DOCX)


# ------------------------------------------------------------------------ normalisation
def test_normalise_text() -> None:
    raw = "Title\r\n\u2022 first\u00a0 item\n\u201cquoted\u201d exam-\nple\x07\n\n\n\nend  "
    assert normalise_text(raw) == 'Title\n- first item\n"quoted" example\n\nend'


def test_estimate_tokens() -> None:
    assert estimate_tokens("") == 0
    assert estimate_tokens("   ") == 0
    assert estimate_tokens("abcd") == 1
    assert estimate_tokens("abcde") == 2


# ------------------------------------------------------------------------------ chunker
DOC = """# Portal SRS

## 1 Approvals

Approvers must approve invoices. Approvers may delegate approvals.
Invoices over 10,000 EUR need a second approval.

## 2 Security

All users must use MFA.
"""


def test_detect_heading_rules() -> None:
    assert detect_heading("## Scope") == (2, "Scope")
    assert detect_heading("3.2 Scheduling and dispatch") == (2, "3.2 Scheduling and dispatch")
    assert detect_heading("FUNCTIONAL REQUIREMENTS") == (1, "FUNCTIONAL REQUIREMENTS")
    assert detect_heading("The system shall export data.") is None
    assert detect_heading("- SSO") is None


def test_chunks_respect_headings_and_record_section_paths() -> None:
    chunks = Chunker(max_tokens=200, overlap_tokens=20).split(DOC)
    assert len(chunks) == 2
    assert chunks[0].section_path == "Portal SRS > 1 Approvals"
    assert chunks[1].section_path == "Portal SRS > 2 Security"
    assert chunks[1].text.startswith("## 2 Security")
    for chunk in chunks:
        assert chunk.text == DOC[chunk.char_start : chunk.char_end]


def test_chunks_obey_budget_and_overlap() -> None:
    sentences = [
        f"Requirement {i} says the dispatcher shall assign job {i} quickly." for i in range(30)
    ]
    text = "## Dispatch\n\n" + " ".join(sentences)
    chunker = Chunker(max_tokens=60, overlap_tokens=20)
    chunks = chunker.split(text)
    assert len(chunks) > 3
    assert all(c.token_count <= 60 for c in chunks)
    assert [c.ordinal for c in chunks] == list(range(len(chunks)))
    for previous, current in pairwise(chunks):
        # Overlap: the next chunk starts inside the previous one, but always makes progress.
        assert current.char_start < previous.char_end
        assert current.char_end > previous.char_end
        overlap = text[current.char_start : previous.char_end]
        assert estimate_tokens(overlap) <= 20
    assert chunks[-1].char_end == len(text)


def test_overlap_can_be_disabled() -> None:
    text = " ".join(f"Sentence number {i} is here." for i in range(40))
    chunks = Chunker(max_tokens=40, overlap_tokens=0).split(text)
    for previous, current in pairwise(chunks):
        assert current.char_start >= previous.char_end


def test_oversized_sentence_is_split_into_word_windows() -> None:
    text = "word " * 600
    chunks = Chunker(max_tokens=50, overlap_tokens=0).split(text.strip())
    assert len(chunks) > 5
    assert all(c.token_count <= 50 for c in chunks)
    assert "".join(c.text.replace(" ", "") for c in chunks) == text.replace(" ", "")


def test_unbroken_token_is_hard_split() -> None:
    chunks = Chunker(max_tokens=30, overlap_tokens=0).split("x" * 1000)
    assert all(c.token_count <= 30 for c in chunks)


def test_list_items_and_abbreviations() -> None:
    text = "- Dr. Who must approve e.g. invoices.\n- Clerks shall export files."
    chunks = Chunker(max_tokens=200, overlap_tokens=0).split(text)
    assert len(chunks) == 1
    assert chunks[0].text == text


def test_empty_text_and_invalid_configuration() -> None:
    assert Chunker().split("") == []
    with pytest.raises(ValueError, match="max_tokens"):
        Chunker(max_tokens=10)
    with pytest.raises(ValueError, match="overlap_tokens"):
        Chunker(max_tokens=100, overlap_tokens=100)
